"""Prepare domestic-sales records and extract Yonyou domestic voucher samples.

The Feishu table uses the literal value ``出口`` for export sales.  For downstream
business routing it is normalized to ``外销`` while the source value is preserved.
Only ``内销`` and ``出口`` are accepted; values such as ``预录`` and ``视同销售``
are reported as excluded.
"""

from __future__ import annotations

import argparse
import collections
import json
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_SOURCE = ROOT / "data/domestic_source_current.json"
DEFAULT_DOMESTIC_OUTPUT = ROOT / "data/domestic_sales_2026.json"
DEFAULT_SAMPLE_OUTPUT = ROOT / "data/domestic_voucher_samples_2026.json"
DEFAULT_WORKBOOK_OUTPUT = ROOT / "vouchers/内销凭证及数据样例_2026.xlsx"
DEFAULT_BUYER_MAPPING = ROOT / "data/ruibei_buyer_mapping_2026.json"

DOMESTIC = "内销"
EXPORT = "出口"
OUTPUT_NAMES = {"内销": "内销", "出口": "外销"}
LINK_FIELD = "合同号_1"
SALES_METHOD_FIELD = "销售方式"

VOUCHER_FIELDS = {
    "id": "凭证ID",
    "displayname": "凭证号",
    "period": "期间",
    "busidate": "业务日期",
    "bodydescription": "摘要",
    "bodies_accsubjectVid_code": "科目编码",
    "bodies_accsubjectVid_name": "科目名称",
    "bodies_currency_name": "币种",
    "bodies_rateOrg": "汇率",
    "debit_original": "原币借方",
    "credit_original": "原币贷方",
    "debit_org": "本币借方",
    "credit_org": "本币贷方",
    "quantity": "数量",
    "price": "单价",
    "auxiliary": "辅助核算",
}

RECORD_COLUMNS = [
    ("合同号_1", "合同号_1"),
    ("合同号", "原合同号"),
    ("出运日期", "出运日期"),
    ("买方英文名", "客户"),
    ("客户名称_脚本用", "内销买方名称"),
    ("客户名称_来源", "买方名称来源"),
    ("睿贝开票申请ID", "开票申请ID"),
    ("纳税识别号（睿贝）", "纳税识别号"),
    ("供应商", "供应商"),
    ("报关品名", "报关品名"),
    ("产品类型", "产品类型"),
    ("报关金额", "含税销售额"),
    ("销项税额", "销项税额"),
    ("不含税人民币收入", "不含税收入"),
    ("不含税成本", "不含税成本"),
    ("采购金额", "采购金额"),
    ("进项税额", "进项税额"),
    ("内陆运费含税", "内陆运费含税"),
    ("报关单号", "报关单号"),
]


def text(value: Any) -> str:
    if isinstance(value, list):
        parts = []
        for item in value:
            parts.append(text(item))
        return "；".join(part for part in parts if part).strip()
    if isinstance(value, dict):
        if value.get("text") is not None:
            return str(value["text"]).strip()
        if value.get("name") is not None:
            return str(value["name"]).strip()
        if value.get("value") is not None:
            return text(value["value"])
        if value.get("id") is not None:
            return str(value["id"]).strip()
        return ""
    return str(value or "").strip()


def number(value: Any) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def read_json(path: Path) -> dict:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def write_json(path: Path, value: Any) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def fetch_source() -> dict:
    sys.path.insert(0, str(ROOT))
    from app.services.feishu_service import list_records

    records = list_records()
    return {
        "fetched_at": datetime.now(timezone.utc).isoformat(),
        "source": {
            "description": "2026报关数据",
            "configuration": "MT_BAO_GUAN_APP_TOKEN / MT_BAO_GUAN_TABLE_ID",
        },
        "records": records,
    }


def normalize_source(source: dict) -> tuple[list[dict], list[dict], collections.Counter]:
    total_methods = collections.Counter()
    normalized = []
    excluded = []
    for wrapped in source.get("records") or []:
        fields = dict(wrapped.get("fields") or wrapped)
        raw_method = text(fields.get(SALES_METHOD_FIELD))
        total_methods[raw_method] += 1
        if raw_method not in OUTPUT_NAMES:
            excluded.append({
                "record_id": wrapped.get("record_id") or wrapped.get("id"),
                "合同号_1": text(fields.get(LINK_FIELD)),
                "销售方式": raw_method,
                "排除原因": "非内销/出口",
            })
            continue

        record = dict(fields)
        record.update({
            "record_id": wrapped.get("record_id") or wrapped.get("id"),
            "销售方式": raw_method,
            "业务链路": OUTPUT_NAMES[raw_method],
            "凭证关联字段": LINK_FIELD,
            "凭证关联值": text(fields.get(LINK_FIELD)),
        })
        normalized.append(record)
    return normalized, excluded, total_methods


def apply_buyer_mapping(records: list[dict], path: Path) -> dict:
    if not path.exists():
        return {"mapping_file": str(path), "updated_records": 0, "updated_contracts": 0}
    mapping = read_json(path)
    by_contract = {
        text(row.get(LINK_FIELD)): row
        for row in mapping.get("results") or []
        if row.get("状态") == "已获取" and text(row.get("公司名称"))
    }
    updated_records = 0
    updated_contracts = set()
    for record in records:
        contract = text(record.get(LINK_FIELD))
        buyer = by_contract.get(contract)
        if not buyer or text(record.get("买方英文名")) or text(record.get("客户名称_脚本用")):
            continue
        record["客户名称_脚本用"] = text(buyer.get("公司名称"))
        source_type = text(buyer.get("来源类型")) or "开票申请"
        if buyer.get("继承自"):
            original_source = text(buyer.get("继承的原单来源类型")) or "开票申请"
            record["客户名称_来源"] = (
                f"睿贝ERP {source_type}（原单 {buyer['继承自']}，{original_source}）"
            )
            record["睿贝原始合同号"] = text(buyer.get("继承自"))
        else:
            record["客户名称_来源"] = "睿贝ERP " + source_type
        applications = buyer.get("开票申请") or buyer.get("继承的开票申请") or []
        record["睿贝开票申请ID"] = applications[0].get("audit_id") if applications else ""
        record["纳税识别号（睿贝）"] = text(buyer.get("纳税识别号"))
        updated_records += 1
        updated_contracts.add(contract)
    return {
        "mapping_file": str(path),
        "mapping_fetched_at": mapping.get("fetched_at"),
        "updated_records": updated_records,
        "updated_contracts": len(updated_contracts),
    }


def group_contracts(records: list[dict]) -> tuple[list[dict], list[dict]]:
    by_contract: dict[str, list[dict]] = collections.defaultdict(list)
    for record in records:
        if record["业务链路"] == "内销":
            by_contract[text(record.get(LINK_FIELD))].append(record)

    groups = []
    anomalies = []
    for contract, rows in sorted(by_contract.items(), key=lambda item: (item[0] == "", item[0])):
        group = {
            "合同号_1": contract,
            "业务链路": "内销",
            "记录数": len(rows),
            "含税销售额": round(sum(number(r.get("报关金额")) for r in rows), 2),
            "不含税收入": round(sum(number(r.get("不含税人民币收入")) for r in rows), 2),
            "销项税额": round(sum(number(r.get("销项税额")) for r in rows), 2),
            "不含税成本": round(sum(number(r.get("不含税成本")) for r in rows), 2),
            "采购金额": round(sum(number(r.get("采购金额")) for r in rows), 2),
            "进项税额": round(sum(number(r.get("进项税额")) for r in rows), 2),
            "客户": sorted({
                text(r.get("买方英文名")) or text(r.get("客户名称_脚本用"))
                for r in rows
                if text(r.get("买方英文名")) or text(r.get("客户名称_脚本用"))
            }),
            "客户名称来源": sorted({
                text(r.get("客户名称_来源"))
                for r in rows
                if text(r.get("客户名称_来源"))
            }),
            "供应商": sorted({text(r.get("供应商")) for r in rows if text(r.get("供应商"))}),
            "报关品名": sorted({text(r.get("报关品名")) for r in rows if text(r.get("报关品名"))}),
            "record_ids": [r.get("record_id") for r in rows],
        }
        groups.append(group)
        if not contract:
            anomalies.append({
                "类型": "合同号为空",
                "合同号_1": "",
                "记录数": len(rows),
                "record_ids": [r.get("record_id") for r in rows],
                "处理": "保留在内销数据中，但不归入可生成凭证组",
            })
    return groups, anomalies


def latest_voucher_file() -> Path | None:
    paths = list((ROOT / "data/yonyou_refresh").glob("*/vouchers/actual_vouchers.json"))
    return max(paths, key=lambda path: path.stat().st_mtime) if paths else None


def group_vouchers(records: list[dict]) -> list[dict]:
    grouped: dict[str, list[dict]] = collections.defaultdict(list)
    for record in records:
        grouped[str(record.get("id") or "")].append(record)
    return [
        {"id": voucher_id, "rows": rows}
        for voucher_id, rows in grouped.items()
        if voucher_id
    ]


def voucher_is_domestic(rows: list[dict]) -> bool:
    return any(
        text(row.get("bodies_accsubjectVid_code")).startswith("500102")
        for row in rows
    )


def match_contract(description: str, contracts: list[str]) -> str:
    # Longest first prevents 25MT-03R371Y from swallowing
    # 25MT-03R371Y-ADD1 when both are present in the domestic source.
    ordered = sorted(contracts, key=len, reverse=True)
    for contract in ordered:
        if contract and description.startswith(contract):
            return contract
    for contract in ordered:
        if contract and contract in description:
            return contract
    return ""


def voucher_sample(voucher: dict, contracts: list[str]) -> dict:
    rows = voucher["rows"]
    descriptions = sorted({text(row.get("bodydescription")) for row in rows})
    matched = sorted({match_contract(desc, contracts) for desc in descriptions} - {""})
    lines = []
    for row in rows:
        lines.append({
            "科目编码": text(row.get("bodies_accsubjectVid_code")),
            "科目名称": text(row.get("bodies_accsubjectVid_name")),
            "摘要": text(row.get("bodydescription")),
            "币种": text(row.get("bodies_currency_name")),
            "汇率": row.get("bodies_rateOrg"),
            "原币借方": number(row.get("debit_original")),
            "原币贷方": number(row.get("credit_original")),
            "本币借方": number(row.get("debit_org")),
            "本币贷方": number(row.get("credit_org")),
            "数量": row.get("quantity"),
            "单价": row.get("price"),
            "辅助核算": text(row.get("auxiliary")),
        })
    return {
        "凭证ID": str(rows[0].get("id") or ""),
        "凭证号": text(rows[0].get("displayname")),
        "期间": text(rows[0].get("period")),
        "业务日期": text(rows[0].get("busidate")),
        "制单人": text(rows[0].get("maker")),
        "审核人": text(rows[0].get("auditor")),
        "摘要": descriptions,
        "匹配合同号_1": matched,
        "类型": "损益结转" if all("损益结转" in desc for desc in descriptions) else "内销业务",
        "行数": len(rows),
        "科目编码": sorted({text(row.get("bodies_accsubjectVid_code")) for row in rows}),
        "分录": lines,
    }


def build_voucher_samples(voucher_data: dict, contracts: list[str]) -> dict:
    vouchers = group_vouchers(voucher_data.get("records") or [])
    domestic = [voucher for voucher in vouchers if voucher_is_domestic(voucher["rows"])]
    samples = [voucher_sample(voucher, contracts) for voucher in domestic]
    samples.sort(key=lambda item: (item["期间"], item["凭证号"], item["凭证ID"]))

    matched_contracts = sorted({
        contract for sample in samples for contract in sample["匹配合同号_1"]
    })
    matched_voucher_ids = {
        sample["凭证ID"] for sample in samples if sample["匹配合同号_1"]
    }
    return {
        "source_file": str(voucher_data.get("_source_file") or ""),
        "fetched_at": voucher_data.get("fetched_at"),
        "scope": voucher_data.get("scope"),
        "record_count": len(voucher_data.get("records") or []),
        "voucher_count": len(vouchers),
        "domestic_voucher_count": len(samples),
        "domestic_line_count": sum(sample["行数"] for sample in samples),
        "matched_domestic_contracts": matched_contracts,
        "matched_voucher_ids": sorted(matched_voucher_ids),
        "samples": samples,
    }


def add_excel_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[Any]]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(headers, start=1):
        values = [str(row[index - 1] or "") for row in rows[:200]]
        width = min(50, max([len(header) * 2] + [len(value) + 2 for value in values]))
        sheet.column_dimensions[get_column_letter(index)].width = max(10, width)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def build_workbook(path: Path, records: list[dict], groups: list[dict], samples: dict,
                   buyer_mapping: dict | None = None) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)

    record_headers = [label for _, label in RECORD_COLUMNS]
    record_rows = [
        [text(record.get(key)) if not isinstance(record.get(key), (int, float)) else record.get(key)
         for key, _ in RECORD_COLUMNS]
        for record in records
        if record["业务链路"] == "内销"
    ]
    add_excel_sheet(workbook, "内销数据", record_headers, record_rows)

    group_headers = [
        "合同号_1", "记录数", "含税销售额", "不含税收入", "销项税额",
        "不含税成本", "采购金额", "进项税额", "客户", "供应商", "报关品名",
    ]
    group_rows = [
        [
            group["合同号_1"], group["记录数"], group["含税销售额"], group["不含税收入"],
            group["销项税额"], group["不含税成本"], group["采购金额"], group["进项税额"],
            "；".join(group["客户"]), "；".join(group["供应商"]), "；".join(group["报关品名"]),
        ]
        for group in groups
    ]
    add_excel_sheet(workbook, "内销分组", group_headers, group_rows)

    voucher_headers = ["凭证ID", "凭证号", "期间", "业务日期", "摘要", "科目编码", "科目名称",
                       "币种", "汇率", "原币借方", "原币贷方", "本币借方", "本币贷方",
                       "数量", "单价", "辅助核算", "匹配合同号_1", "类型"]
    voucher_rows = []
    for sample in samples["samples"]:
        for line in sample["分录"]:
            voucher_rows.append([
                sample["凭证ID"], sample["凭证号"], sample["期间"], sample["业务日期"],
                line["摘要"], line["科目编码"], line["科目名称"], line["币种"], line["汇率"],
                line["原币借方"], line["原币贷方"], line["本币借方"], line["本币贷方"],
                line["数量"], line["单价"], line["辅助核算"],
                "；".join(sample["匹配合同号_1"]), sample["类型"],
            ])
    add_excel_sheet(workbook, "用友内销凭证", voucher_headers, voucher_rows)

    index_headers = ["凭证ID", "凭证号", "期间", "业务日期", "摘要", "行数", "科目编码",
                     "匹配合同号_1", "类型"]
    index_rows = [
        [
            sample["凭证ID"], sample["凭证号"], sample["期间"], sample["业务日期"],
            "；".join(sample["摘要"]), sample["行数"], "、".join(sample["科目编码"]),
            "；".join(sample["匹配合同号_1"]), sample["类型"],
        ]
        for sample in samples["samples"]
    ]
    add_excel_sheet(workbook, "凭证索引", index_headers, index_rows)

    if buyer_mapping:
        buyer_headers = [
            "合同号_1", "状态", "公司名称", "纳税识别号", "开票申请数",
            "搜索值", "开票申请ID", "相关单据", "来源类型", "继承自",
            "来源字段",
        ]
        buyer_rows = []
        for result in buyer_mapping.get("results") or []:
            applications = result.get("开票申请") or []
            buyer_rows.append([
                result.get("合同号_1"),
                result.get("状态"),
                result.get("公司名称"),
                result.get("纳税识别号"),
                result.get("开票申请数"),
                "；".join(result.get("搜索值") or []),
                "；".join(str(item.get("audit_id") or "") for item in applications),
                "；".join(str(item.get("related_bill") or "") for item in applications),
                result.get("来源类型") or "开票申请",
                result.get("继承自") or "",
                (result.get("来源单据") or {}).get("来源字段") or "",
            ])
        add_excel_sheet(workbook, "睿贝买方补全", buyer_headers, buyer_rows)

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--fetch", action="store_true", help="Fetch a fresh source snapshot before preparing")
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--vouchers", type=Path, help="Defaults to the newest Yonyou refresh")
    parser.add_argument("--output", type=Path, default=DEFAULT_DOMESTIC_OUTPUT)
    parser.add_argument("--sample-output", type=Path, default=DEFAULT_SAMPLE_OUTPUT)
    parser.add_argument("--workbook-output", type=Path, default=DEFAULT_WORKBOOK_OUTPUT)
    parser.add_argument("--buyer-mapping", type=Path, default=DEFAULT_BUYER_MAPPING)
    args = parser.parse_args()

    source = fetch_source() if args.fetch else read_json(args.source)
    records, excluded, methods = normalize_source(source)
    buyer_mapping = read_json(args.buyer_mapping) if args.buyer_mapping.exists() else None
    buyer_enrichment = apply_buyer_mapping(records, args.buyer_mapping)
    groups, anomalies = group_contracts(records)
    domestic_records = [record for record in records if record["业务链路"] == "内销"]
    export_records = [record for record in records if record["业务链路"] == "外销"]

    domestic_output = {
        "fetched_at": source.get("fetched_at"),
        "source": source.get("source"),
        "filter_rule": {
            "include_source_values": ["内销", "出口"],
            "business_mapping": OUTPUT_NAMES,
            "exclude_other_values": True,
            "domestic_voucher_unit": LINK_FIELD,
            "customs_declaration_link": "出口凭证按报关单号；内销没有报关单号",
        },
        "total_records": len(source.get("records") or []),
        "source_method_counts": dict(methods),
        "buyer_enrichment": buyer_enrichment,
        "domestic_record_count": len(domestic_records),
        "domestic_contract_count": sum(1 for group in groups if group["合同号_1"]),
        "domestic_contract_record_count": sum(
            group["记录数"] for group in groups if group["合同号_1"]
        ),
        "export_record_count": len(export_records),
        "excluded_record_count": len(excluded),
        "excluded": excluded,
        "anomalies": anomalies,
        "groups": groups,
        "records": domestic_records,
    }
    write_json(args.output, domestic_output)

    voucher_path = args.vouchers or latest_voucher_file()
    if voucher_path:
        voucher_data = read_json(voucher_path)
        voucher_data["_source_file"] = str(voucher_path.resolve())
        contracts = [group["合同号_1"] for group in groups if group["合同号_1"]]
        samples = build_voucher_samples(voucher_data, contracts)
        source_contracts = set(contracts)
        matched_contracts = set(samples["matched_domestic_contracts"])
        samples["unmatched_domestic_contracts"] = sorted(source_contracts - matched_contracts)
        samples["unmatched_contract_count"] = len(samples["unmatched_domestic_contracts"])
        write_json(args.sample_output, samples)
        build_workbook(args.workbook_output, records, groups, samples, buyer_mapping)
        sample_summary = {
            "source": str(voucher_path),
            "domestic_vouchers": samples["domestic_voucher_count"],
            "domestic_lines": samples["domestic_line_count"],
            "matched_contracts": len(matched_contracts),
            "unmatched_contracts": samples["unmatched_contract_count"],
        }
    else:
        sample_summary = {"source": None, "reason": "No Yonyou voucher refresh found"}

    print(json.dumps({
        "feishu": {
            "total": domestic_output["total_records"],
            "methods": domestic_output["source_method_counts"],
            "domestic_records": domestic_output["domestic_record_count"],
            "domestic_contracts": domestic_output["domestic_contract_count"],
            "domestic_contract_records": domestic_output["domestic_contract_record_count"],
            "excluded": domestic_output["excluded_record_count"],
            "anomalies": len(anomalies),
        },
        "yonyou": sample_summary,
        "buyer_enrichment": buyer_enrichment,
        "outputs": {
            "domestic": str(args.output.resolve()),
            "samples": str(args.sample_output.resolve()) if voucher_path else None,
            "workbook": str(args.workbook_output.resolve()) if voucher_path else None,
        },
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
