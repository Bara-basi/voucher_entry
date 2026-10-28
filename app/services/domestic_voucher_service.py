"""Domestic vouchers from the confirmed template; never reads historical fitted rules."""
from __future__ import annotations

import collections
import copy
import datetime as dt
import json
import re
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from app.settings import voucher_maker
from . import voucher_service as legacy
from . import ruibei_order_cache as orders

ROOT = Path(__file__).resolve().parents[2]
RUI_BEI_BUYER_MAPPING = ROOT / ".chche/domestic_orders.json"
CONTRACT_FIELD = "合同号_1"
SOURCE_METHOD = "内销"
TEMPLATE_VERSION = "domestic-template-20261027-export-customer-policy-v4"
FEE_FIELDS = ("运费", "保费", "港杂费", "港杂费用")
text = legacy.to_text
normalize_name = legacy.normalize_name


class DomesticVoucherError(Exception):
    pass


def number(value, default=Decimal("0")):
    if isinstance(value, dict):
        return number(value.get("value", value.get("text")), default)
    if isinstance(value, list):
        return number(value[0], default) if len(value) == 1 else default
    if value is None or value == "" or isinstance(value, bool):
        return default
    try:
        result = Decimal(str(value).strip())
        return result if result.is_finite() and abs(result) < Decimal("1e14") else default
    except (InvalidOperation, ValueError, TypeError):
        return default


def money(value):
    return value.quantize(Decimal(".01"), rounding=ROUND_HALF_UP) if value is not None else None


def enrich_buyer_names(records, masters=None):
    """Apply order-scoped customerName values, preserving untouched Feishu source."""
    result = copy.deepcopy(records)
    cache = orders.load_cache(RUI_BEI_BUYER_MAPPING)
    for record in result:
        if text(record.get("销售方式")) != SOURCE_METHOD:
            continue
        record.pop("_erp_customer_conflict", None)
        if record.get("_manual_customer_name"):
            continue
        origin = text(record.get("客户名称_来源"))
        if ("睿贝ERP" in origin and not any(module in origin for module in ("内销订单", "外销订单"))) or any(word in origin for word in ("发货审批", "收货单位")):
            record["客户名称_脚本用"] = ""
            record["客户名称_来源"] = ""
        item = orders.resolve(cache, text(record.get(CONTRACT_FIELD)))
        if item["状态"] != "已获取":
            if item["订单"]:
                record["_erp_customer_conflict"] = True
            continue
        record["客户名称_脚本用"] = item["公司名称"]
        record["客户名称_来源"] = "睿贝ERP" + item["来源类型"].removesuffix("客户名称") + (" / ADD继承" if item["继承自"] else "")
        record["睿贝内销订单ID"] = ",".join(row["order_id"] for row in item["订单"])
        record.pop("睿贝原始合同号", None)
        record.pop("纳税识别号（睿贝）", None)
        if item["继承自"]:
            record["睿贝原始合同号"] = item["继承自"]
    return result


def customer_for_records(records, masters):
    if any(record.get("_erp_customer_conflict") for record in records):
        return ""
    return masters.customer_for_records(records) or ""


def customer_failure_reason(records, masters):
    cache = orders.load_cache(RUI_BEI_BUYER_MAPPING)
    messages = []
    codes = {masters.customer_for_record(record) for record in records}
    if len(codes - {None, ""}) > 1:
        return "同一合同存在不同客户，请核实实际客户主体"
    for record in records:
        if masters.customer_for_record(record) and not record.get("_erp_customer_conflict"):
            continue
        name = text(record.get("客户名称_脚本用")) or text(record.get("买方英文名"))
        status = orders.resolve(cache, text(record.get(CONTRACT_FIELD)))["状态"]
        if status == "内销订单客户冲突":
            message = "睿贝订单存在不同客户，需人工核对，未自动选择"
        elif status == "内销订单客户字段为空":
            message = "已找到睿贝订单，但客户名称字段为空"
        elif text(record.get("客户名称_来源")).startswith(("睿贝ERP内销订单", "睿贝ERP外销订单")):
            message = f"睿贝订单已取得客户“{name}”，但未匹配用友客户档案，请维护映射"
        elif cache:
            message = "睿贝内销及外销订单均未找到本合同" + (f"，原客户名“{name}”也无法匹配用友档案" if name else "")
        else:
            message = "尚无完整睿贝订单缓存，请先获取订单客户信息"
        if message not in messages:
            messages.append(message)
    return "；".join(messages) or "客户无法唯一匹配"


def _row(row_id, summary, subject, amount, *, debit, customer="", supplier="", inventory="", quantity=None, unit_price=None):
    value = str(money(amount)) if amount is not None else ""
    return {"id": row_id, "摘要": summary, "科目编码": subject, "币种名称": "CNY", "汇率": "1",
            "原币借方": value if debit else "0", "原币贷方": "0" if debit else value,
            "借方金额": value if debit else "0", "贷方金额": "0" if debit else value,
            "客户编码": customer, "供应商编码": supplier, "存货档案": inventory,
            "借方数量": str(quantity) if quantity is not None and debit else "",
            "贷方数量": str(quantity) if quantity is not None and not debit else "",
            "单价": str(money(unit_price)) if unit_price is not None else ""}


def build_domestic_voucher(records, masters=None):
    if not records:
        raise DomesticVoucherError("空合同分组")
    masters = masters or legacy.MasterData()
    records = enrich_buyer_names(records, masters)
    contract = text(records[0].get(CONTRACT_FIELD))
    if not contract:
        raise DomesticVoucherError("合同号_1 为空")
    if any(text(r.get(CONTRACT_FIELD)) != contract for r in records):
        raise DomesticVoucherError("同一凭证不能混入不同合同")
    issues, warnings = [], []

    def issue(message, row_id="", field="", record=None, source_field="", severity="error"):
        issues.append({"id": f"domestic:{row_id}:{field}:{len(issues)}", "message": message,
                       "severity": severity, "row_id": row_id, "field": field,
                       "record_id": text((record or {}).get("record_id")), "source_field": source_field})

    def amount(record, field, row_id, target):
        value = number(record.get(field), None)
        if value is None:
            issue(f"缺少或无效的{field}，请补充来源后重新计算", row_id, target, record, field)
        return value

    def total(field, row_id, target):
        values = [amount(r, field, row_id, target) for r in records]
        return money(sum(values, Decimal("0"))) if all(v is not None for v in values) else None

    customer = customer_for_records(records, masters)
    if not customer:
        issue("未匹配客户：" + customer_failure_reason(records, masters),
              "income:ar", "客户编码", records[0], "客户名称_脚本用")
    gross = total("报关金额", "income:ar", "借方金额")
    revenue = total("不含税人民币收入", "income:revenue", "贷方金额")
    tax = total("销项税额", "income:tax", "贷方金额")
    if None not in (gross, revenue, tax) and gross != revenue + tax:
        issue(f"收入区金额不平：应收={gross}，收入+销项税={money(revenue + tax)}，请核实来源",
              record=records[0], source_field="报关金额")
    summary = contract + "内销收入"
    rows = [_row("income:ar", summary, "112201", gross, debit=True, customer=customer),
            _row("income:revenue", summary, "50010201", revenue, debit=False),
            _row("income:tax", summary, "22210102", tax, debit=False)]
    for index, record in enumerate(records, 1):
        key = text(record.get("record_id")) or f"record-{index}"
        if text(record.get("销售方式")) != SOURCE_METHOD:
            issue("销售方式不是内销，请核实来源", record=record, source_field="销售方式")
        if text(record.get("币种")).upper() not in ("CNY", "RMB", "人民币"):
            issue("币种缺失或非人民币内销：请确认处理规则后再导出", record=record, source_field="币种")
        for field in FEE_FIELDS:
            if text(record.get(field)):
                fee = number(record.get(field), None)
                if fee is None or fee != 0:
                    issue(f"内销存在非零或无效{field}：请确认费用规则后再导出", record=record, source_field=field)
        raw_cost = amount(record, "不含税成本", key + ":stock", "借方金额")
        cost = money(raw_cost)
        itax = money(amount(record, "进项税额", key + ":tax", "借方金额"))
        purchase = money(amount(record, "采购金额", key + ":ap", "贷方金额"))
        if None not in (cost, itax, purchase) and purchase != cost + itax:
            issue(f"采购金额与不含税成本及进项税额不一致：应付={purchase}，成本+税={money(cost + itax)}",
                  record=record, source_field="采购金额")
        supplier_raw = text(record.get("供应商"))
        supplier = "" if re.search(r"[,，、;；]", supplier_raw) else (
            masters.supplier_by_name.get(normalize_name(supplier_raw))
            or masters.supplier_code(record.get("供应商简称"), supplier_raw) or "")
        inventory = masters.inventory_code(text(record.get("报关品名"))) or ""
        subject = legacy.INVENTORY_SUBJECT.get(legacy.PRODUCT_CODE_MAPPING.get(inventory[:4]), "")
        weight = number(record.get("报关重量"), None)
        if weight is None or weight <= 0:
            issue("报关重量缺失、无效或不大于零，请核实数量；不使用开票数量替代",
                  record=record, source_field="报关重量")
            weight = None
        unit_price = None
        if raw_cost is not None and weight:
            try:
                unit_price = money(raw_cost / weight)
            except InvalidOperation:
                issue("成本或报关重量导致单价超出数值范围，请核实来源", record=record, source_field="报关重量")
        if not supplier:
            issue("供应商缺失、多值或无法匹配，请补充来源或人工选择供应商", key + ":ap", "供应商编码", record, "供应商")
        if not inventory or not subject:
            issue("未匹配存货或库存科目，请维护产品映射", key + ":stock", "存货档案", record, "报关品名")
        purchase_summary, cost_summary = contract + "采购商品成本", contract + "结转采购成本"
        rows.extend([
            _row(key + ":stock", purchase_summary, subject, cost, debit=True,
                 inventory=inventory, quantity=weight, unit_price=unit_price),
            _row(key + ":tax", purchase_summary + "-专票", "22210101", itax, debit=True),
            _row(key + ":ap", purchase_summary, "220201", purchase, debit=False, supplier=supplier),
            _row(key + ":cost", cost_summary, "540102", cost, debit=True),
            _row(key + ":out", cost_summary, subject, cost, debit=False,
                 inventory=inventory, quantity=weight, unit_price=unit_price),
        ])
    debit = money(sum((number(r["借方金额"]) for r in rows), Decimal("0")))
    credit = money(sum((number(r["贷方金额"]) for r in rows), Decimal("0")))
    if debit != credit:
        issue(f"凭证借贷不平衡：借方={debit}，贷方={credit}，差额={money(debit-credit)}")
    return {"合同号_1": contract, "rows": rows, "issues": issues, "warnings": warnings,
            "record_ids": [r.get("record_id") for r in records], "effective_source": records,
            "template_version": TEMPLATE_VERSION,
            "totals": {"debit": str(debit), "credit": str(credit), "difference": str(money(debit-credit))}}


def generate_all(payload, masters=None):
    masters = masters or legacy.MasterData()
    groups = collections.defaultdict(list)
    for item in payload.get("records", []):
        record = dict(item.get("fields", item))
        record["record_id"] = item.get("record_id") or item.get("id") or record.get("record_id")
        if text(record.get("销售方式")) == SOURCE_METHOD:
            groups[text(record.get(CONTRACT_FIELD))].append(record)
    generated, failed, skipped = [], [], []
    for contract, records in sorted(groups.items()):
        if not contract:
            skipped.append({"合同号_1": "", "原因": "合同号_1 为空", "记录数": len(records)})
            continue
        try:
            generated.append(build_domestic_voucher(records, masters))
        except DomesticVoucherError as error:
            failed.append({"合同号_1": contract, "原因": str(error)})
    return {"generated": generated, "failed": failed, "skipped": skipped,
            "generated_count": len(generated), "failed_count": len(failed), "skipped_count": len(skipped),
            "template_version": TEMPLATE_VERSION}


def _excel_row(row):
    fields = {"汇率", "单价", "借方数量", "贷方数量", "原币借方", "原币贷方", "借方金额", "贷方金额"}
    return {k: (float(number(v)) if v != "" else None) if k in fields else v for k, v in row.items()}


def write_excel(result: dict, path: Path, posting_date: dt.date,
                maker: str | None = None) -> None:
    maker = maker or voucher_maker()
    book, sheet, columns = legacy._new_sheet(path)
    sheet._cell_overwrite_ok = True
    row_number = 1
    for voucher_id, voucher in enumerate([v for v in result["generated"] if not any(i["severity"] == "error" for i in v.get("issues", []))], start=1):
        for row in voucher["rows"]:
            legacy._write_row(sheet, columns, row_number, _excel_row(row), voucher_id, posting_date)
            for field, value in (("制单人", maker), ("凭证类别", "记")):
                if field in columns:
                    sheet.write(row_number, columns[field], value)
            row_number += 1
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".part")
    book.save(str(temporary))
    temporary.replace(path)



def write_single_excels(result: dict, directory: Path, posting_date: dt.date,
                        maker: str | None = None) -> Path:
    maker = maker or voucher_maker()
    directory.mkdir(parents=True, exist_ok=True)
    index_rows = []
    for index, voucher in enumerate([v for v in result["generated"] if not any(i["severity"] == "error" for i in v.get("issues", []))], start=1):
        contract = voucher["合同号_1"]
        safe_contract = re.sub(r"[^0-9A-Za-z\u4e00-\u9fa5&_-]", "_", contract)
        filename = f"{index:03d}_{safe_contract}.xls"
        path = directory / filename
        book, sheet, columns = legacy._new_sheet(path)
        sheet._cell_overwrite_ok = True
        for row_number, row in enumerate(voucher["rows"], start=1):
            legacy._write_row(sheet, columns, row_number, _excel_row(row), 1, posting_date)
            for field, value in (("制单人", maker), ("凭证类别", "记")):
                if field in columns:
                    sheet.write(row_number, columns[field], value)
        temporary = path.with_suffix(".part")
        book.save(str(temporary))
        temporary.replace(path)
        voucher["file"] = filename
        index_rows.append([
            index,
            contract,
            filename,
            len(voucher["rows"]),
            voucher.get("template_version", ""),
            len(voucher.get("issues", [])),
            len(voucher["warnings"]),
            "；".join(sorted({warning["原因"] for warning in voucher["warnings"]})),
        ])

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "单张凭证索引"
    headers = [
        "序号", "合同号_1", "导入文件", "分录行数",
        "模板版本", "异常数", "警告数", "警告原因",
    ]
    sheet.append(headers)
    for row in index_rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(headers, start=1):
        widths = [len(header) * 2]
        widths.extend(min(50, len(str(row[index - 1] or "")) + 2) for row in index_rows)
        sheet.column_dimensions[get_column_letter(index)].width = max(12, max(widths))
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    index_path = directory / "单张凭证索引.xlsx"
    workbook.save(index_path)
    return index_path



def save_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)
