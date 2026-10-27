"""Independent historical comparisons; differences are review evidence, not ground truth.

This module never participates in building workbench vouchers or choosing accounts.
"""
from __future__ import annotations

import collections
import re
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

def text(value: Any) -> str:
    if isinstance(value, list):
        return "；".join(
            text(item)
            for item in value
        ).strip("； ")
    if isinstance(value, dict):
        if value.get("text") is not None:
            return str(value["text"]).strip()
        if value.get("name") is not None:
            return str(value["name"]).strip()
        if value.get("value") is not None:
            return text(value["value"])
        return ""
    return str(value or "").strip()



def number(value: Any, default: Decimal | None = Decimal("0")) -> Decimal | None:
    if value in (None, ""):
        return default
    try:
        parsed = Decimal(str(value))
        return parsed if parsed.is_finite() else default
    except (InvalidOperation, ValueError, TypeError):
        return default



def money(value: Any) -> Decimal:
    return (number(value) or Decimal("0")).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)



def aux_value(line: dict, label: str) -> dict:
    match = re.search(
        rf"【{re.escape(label)}:([^/】]+)/([^】]+)】",
        text(line.get("辅助核算")),
    )
    return {"code": match.group(1), "name": match.group(2)} if match else {}



def evaluate(generated: dict, sample_payload: dict) -> dict:
    """Independent post-generation comparison; never called by generation."""
    sample_by_contract = collections.defaultdict(list)
    for sample in sample_payload.get("samples") or []:
        if sample.get("类型") != "内销业务":
            continue
        for contract in sample.get("匹配合同号_1") or []:
            sample_by_contract[contract].append(sample)

    comparisons = []
    skipped_without_contract_lines = []
    suffixes = sorted({
        "内销收入",
        "销收入",
        "内采购商品收入-专票",
        "内采购商品收入",
        "内销采购商品收入-专票",
        "内销采购商品收入",
        "采购商品成本-专票",
        "采购商品成本",
        "结转内销采购成本",
        "结转采购成本",
    }, key=len, reverse=True)

    def summary_contract(description: str) -> str:
        value = description.replace("\xa0", " ").strip()
        if value.startswith("暂估-"):
            value = value[len("暂估-"):]
        for suffix in suffixes:
            if value.endswith(suffix):
                return value[:-len(suffix)].strip()
        return ""

    for voucher in generated.get("generated") or []:
        contract = voucher["合同号_1"]
        samples = sample_by_contract.get(contract, [])
        if not samples:
            continue
        generated_revenue = money(sum(
            number(row["贷方金额"]) or Decimal("0")
            for row in voucher["rows"]
            if text(row["科目编码"]).startswith("500102")
        ))
        sample_lines = []
        for sample in samples:
            for line in sample.get("分录") or []:
                description = text(line.get("摘要"))
                if summary_contract(description) == contract:
                    sample_lines.append(line)
        if not sample_lines:
            skipped_without_contract_lines.append(contract)
            continue
        generated_gross = money(sum(
            number(row["借方金额"]) or Decimal("0")
            for row in voucher["rows"]
            if text(row["摘要"]).endswith("内销收入")
        ))
        sample_gross = money(sum(
            number(line.get("本币借方")) or Decimal("0")
            for line in sample_lines
            if text(line.get("摘要")).endswith("内销收入")
        ))
        generated_tax = money(sum(
            number(row["贷方金额"]) or Decimal("0")
            for row in voucher["rows"]
            if text(row["科目编码"]).startswith("22210102")
        ))
        sample_tax = money(sum(
            number(line.get("本币贷方")) or Decimal("0")
            for line in sample_lines
            if text(line.get("科目编码")).startswith("22210102")
        ))
        sample_revenue = money(sum(
            number(line.get("本币贷方")) or Decimal("0")
            for line in sample_lines
            if text(line.get("科目编码")).startswith("500102")
        ))
        comparisons.append({
            "合同号_1": contract,
            "导入文件": voucher.get("file", ""),
            "生成不含税收入": str(generated_revenue),
            "历史不含税收入": str(sample_revenue),
            "差异": str(generated_revenue - sample_revenue),
            "生成含税收入": str(generated_gross),
            "历史含税收入": str(sample_gross),
            "生成销项税": str(generated_tax),
            "历史销项税": str(sample_tax),
            "收入一致": abs(generated_revenue - sample_revenue) <= Decimal("0.02"),
            "含税收入一致": abs(generated_gross - sample_gross) <= Decimal("0.02"),
            "销项税一致": abs(generated_tax - sample_tax) <= Decimal("0.02"),
            "一致": (
                abs(generated_revenue - sample_revenue) <= Decimal("0.02")
                and abs(generated_gross - sample_gross) <= Decimal("0.02")
                and abs(generated_tax - sample_tax) <= Decimal("0.02")
            ),
        })
    return {
        "compared_contracts": len(comparisons),
        "exact_revenue_matches": sum(item["收入一致"] for item in comparisons),
        "exact_sales_matches": sum(item["一致"] for item in comparisons),
        "skipped_without_contract_lines": skipped_without_contract_lines,
        "comparisons": comparisons,
    }



def evaluate_full_lines(generated: dict, sample_payload: dict,
                        voucher_records: list[dict] | None = None) -> dict:
    """Compare every accounting field of comparable generated/manual vouchers."""
    suffixes = sorted({
        "内销收入",
        "销收入",
        "内采购商品收入-专票",
        "内采购商品收入",
        "内销采购商品收入-专票",
        "内销采购商品收入",
        "采购商品成本-专票",
        "采购商品成本",
        "结转内销采购成本",
        "结转采购成本",
    }, key=len, reverse=True)

    def summary_contract(value: str) -> str:
        description = text(value).replace("\xa0", " ").strip()
        if description.startswith("暂估-"):
            description = description[len("暂估-"):]
        for suffix in suffixes:
            if description.endswith(suffix):
                return description[:-len(suffix)].strip()
        return ""

    def canonical(line: dict, sample: bool) -> dict:
        if sample:
            if line.get("科目编码") or line.get("辅助核算") is not None:
                debit = number(line.get("本币借方"), Decimal("0")) or Decimal("0")
                credit = number(line.get("本币贷方"), Decimal("0")) or Decimal("0")
                customer = aux_value(line, "客户").get("code", "")
                supplier = aux_value(line, "供应商").get("code", "")
                inventory = aux_value(line, "存货档案").get("code", "")
                quantity = line.get("数量")
                unit_price = line.get("单价")
                summary = text(line.get("摘要"))
                code = text(line.get("科目编码"))
            else:
                debit = number(line.get("debit_org"), Decimal("0")) or Decimal("0")
                credit = number(line.get("credit_org"), Decimal("0")) or Decimal("0")
                auxiliaries = {
                    text(item.get("fieldname")): text((item.get("value") or {}).get("code"))
                    for item in line.get("clientAuxiliaryDTOS") or []
                }
                customer = auxiliaries.get("客户", "")
                supplier = auxiliaries.get("供应商", "")
                inventory = auxiliaries.get("存货档案", "")
                quantity = line.get("quantity")
                unit_price = line.get("price")
                summary = text(line.get("bodydescription"))
                code = text(line.get("bodies_accsubjectVid_code"))
        else:
            code = text(line.get("科目编码"))
            debit = number(line.get("借方金额"), Decimal("0")) or Decimal("0")
            credit = number(line.get("贷方金额"), Decimal("0")) or Decimal("0")
            customer = text(line.get("客户编码"))
            supplier = text(line.get("供应商编码"))
            inventory = text(line.get("存货档案"))
            quantity = line.get("借方数量") or line.get("贷方数量")
            unit_price = line.get("单价")
            summary = text(line.get("摘要"))
        quantity_value = number(quantity, None)
        return {
            "科目编码": code,
            "借方": str(money(debit)),
            "贷方": str(money(credit)),
            "客户编码": customer,
            "供应商编码": supplier,
            "存货档案": inventory,
            "数量": str(quantity_value.normalize()) if quantity_value is not None else "",
            "单价": str(money(unit_price)) if unit_price not in (None, "") else "",
            "摘要": summary,
        }

    def without_summary(line: dict) -> tuple:
        return tuple((key, line[key]) for key in line if key != "摘要")

    def relaxed_line(line: dict) -> tuple:
        values = dict(without_summary(line))
        if values["科目编码"] in ("22210101", "22210107"):
            values["科目编码"] = "22210101/22210107"
        return tuple(sorted(values.items()))

    def account_totals(lines: list[dict]) -> dict:
        totals = collections.defaultdict(lambda: {"借方": Decimal("0"), "贷方": Decimal("0")})
        for line in lines:
            totals[line["科目编码"]]["借方"] += Decimal(line["借方"])
            totals[line["科目编码"]]["贷方"] += Decimal(line["贷方"])
        return {
            account: {
                "借方": str(money(values["借方"])),
                "贷方": str(money(values["贷方"])),
            }
            for account, values in totals.items()
        }

    def dimension_totals(lines: list[dict], include_aux: bool) -> dict:
        totals = collections.defaultdict(Decimal)
        for line in lines:
            debit = Decimal(line["借方"])
            credit = Decimal(line["贷方"])
            account = line["科目编码"]
            if account in ("22210101", "22210107"):
                account = "22210101/22210107"
            prefixes = [account, "借" if debit else "贷"]
            if include_aux:
                prefixes.extend([
                    line["客户编码"], line["供应商编码"], line["存货档案"],
                ])
            totals[tuple(prefixes)] += debit or credit
        return dict(totals)

    sample_by_contract = collections.defaultdict(list)
    voucher_ids_by_contract = collections.defaultdict(set)
    if voucher_records:
        for record in voucher_records:
            contract = summary_contract(record.get("bodydescription"))
            if contract:
                sample_by_contract[contract].append(record)
                voucher_ids_by_contract[contract].add(text(record.get("id")))
    else:
        for sample in sample_payload.get("samples") or []:
            if sample.get("类型") != "内销业务":
                continue
            for line in sample.get("分录") or []:
                contract = summary_contract(line.get("摘要"))
                if contract:
                    sample_by_contract[contract].append(line)
                    voucher_ids_by_contract[contract].add(text(sample.get("凭证ID")))

    comparisons = []
    for voucher in generated.get("generated") or []:
        contract = voucher["合同号_1"]
        historical = sample_by_contract.get(contract, [])
        if not historical:
            continue
        actual = [canonical(line, False) for line in voucher["rows"]]
        expected = [canonical(line, True) for line in historical]
        actual_without_summary = collections.Counter(without_summary(line) for line in actual)
        expected_without_summary = collections.Counter(without_summary(line) for line in expected)
        actual_relaxed = collections.Counter(relaxed_line(line) for line in actual)
        expected_relaxed = collections.Counter(relaxed_line(line) for line in expected)
        actual_full = collections.Counter(tuple(line.items()) for line in actual)
        expected_full = collections.Counter(tuple(line.items()) for line in expected)
        actual_accounts = account_totals(actual)
        expected_accounts = account_totals(expected)
        actual_dimensions = dimension_totals(actual, True)
        expected_dimensions = dimension_totals(expected, True)
        actual_account_groups = dimension_totals(actual, False)
        expected_account_groups = dimension_totals(expected, False)
        exact_common = sum((actual_without_summary & expected_without_summary).values())
        dimension_common = sum(
            1
            for key, value in actual_dimensions.items()
            if expected_dimensions.get(key) == value
        )
        account_common = sum(
            1
            for key, value in actual_account_groups.items()
            if expected_account_groups.get(key) == value
        )
        account_differences = []
        for account in sorted(set(actual_accounts) | set(expected_accounts)):
            generated_value = actual_accounts.get(account, {"借方": "0.00", "贷方": "0.00"})
            expected_value = expected_accounts.get(account, {"借方": "0.00", "贷方": "0.00"})
            if generated_value != expected_value:
                account_differences.append({
                    "科目编码": account,
                    "生成借方": generated_value["借方"],
                    "历史借方": expected_value["借方"],
                    "生成贷方": generated_value["贷方"],
                    "历史贷方": expected_value["贷方"],
                })
        relaxed = actual_relaxed == expected_relaxed
        reasons = []
        if not relaxed:
            if len(actual) != len(expected):
                reasons.append(f"分录数不同：生成 {len(actual)} 行，历史 {len(expected)} 行")
            actual_core = collections.Counter(
                (line["科目编码"], line["借方"], line["贷方"])
                if line["科目编码"] not in ("22210101", "22210107")
                else ("22210101/22210107", line["借方"], line["贷方"])
                for line in actual
            )
            expected_core = collections.Counter(
                (line["科目编码"], line["借方"], line["贷方"])
                if line["科目编码"] not in ("22210101", "22210107")
                else ("22210101/22210107", line["借方"], line["贷方"])
                for line in expected
            )
            if actual_core == expected_core:
                reasons.append("科目金额相同，但辅助核算、数量或单价不同")
            else:
                actual_accounts_relaxed = collections.defaultdict(Decimal)
                expected_accounts_relaxed = collections.defaultdict(Decimal)
                for line in actual:
                    account = (
                        "22210101/22210107"
                        if line["科目编码"] in ("22210101", "22210107")
                        else line["科目编码"]
                    )
                    actual_accounts_relaxed[account] += Decimal(line["借方"]) - Decimal(line["贷方"])
                for line in expected:
                    account = (
                        "22210101/22210107"
                        if line["科目编码"] in ("22210101", "22210107")
                        else line["科目编码"]
                    )
                    expected_accounts_relaxed[account] += Decimal(line["借方"]) - Decimal(line["贷方"])
                differing = [
                    account for account in sorted(set(actual_accounts_relaxed) | set(expected_accounts_relaxed))
                    if actual_accounts_relaxed[account] != expected_accounts_relaxed[account]
                ]
                if differing:
                    reasons.append("科目净额不同：" + "、".join(differing))
                else:
                    reasons.append("借方/贷方拆分不同")
        comparisons.append({
            "合同号_1": contract,
            "生成行数": len(actual),
            "历史行数": len(expected),
            "历史凭证ID": sorted(voucher_ids_by_contract[contract]),
            "会计字段一致": actual_without_summary == expected_without_summary,
            "忽略顺序及暂估进项科目后一致": relaxed,
            "剩余不一致原因": reasons,
            "科目汇总一致": not account_differences,
            "科目汇总差异": account_differences,
            "行顺序一致": actual == expected,
            "包含摘要完全一致": actual_full == expected_full,
            "完全一致分录数": exact_common,
            "同科目借方/贷方及辅助核算汇总匹配数": dimension_common,
            "同科目借贷总额匹配数": account_common,
            "存在可匹配分录": bool(exact_common or dimension_common or account_common),
            "生成多出行": [
                dict(line) for line, count in (actual_without_summary - expected_without_summary).items()
                for _ in range(count)
            ],
            "历史多出行": [
                dict(line) for line, count in (expected_without_summary - actual_without_summary).items()
                for _ in range(count)
            ],
        })
    generated_contracts = {
        item["合同号_1"] for item in generated.get("generated") or []
    }
    not_made_contracts = sorted(generated_contracts - set(sample_by_contract))
    return {
        "compared_contracts": len(comparisons),
        "exact_accounting_matches": sum(item["会计字段一致"] for item in comparisons),
        "relaxed_matches": sum(item["忽略顺序及暂估进项科目后一致"] for item in comparisons),
        "exact_account_total_matches": sum(item["科目汇总一致"] for item in comparisons),
        "exact_ordered_matches": sum(item["行顺序一致"] for item in comparisons),
        "exact_full_matches": sum(item["包含摘要完全一致"] for item in comparisons),
        "loose_matched_contracts": sum(item["存在可匹配分录"] for item in comparisons),
        "loose_mismatched_contracts": sum(
            not item["存在可匹配分录"] for item in comparisons
        ),
        "not_made_or_not_found_contracts": not_made_contracts,
        "comparisons": comparisons,
    }



def write_reconciliation_workbook(report: dict, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    full = report["full_line_comparison"]
    workbook = Workbook()
    workbook.remove(workbook.active)

    def add_sheet(title: str, headers: list[str], rows: list[list[Any]]) -> None:
        target = workbook.create_sheet(title)
        target.append(headers)
        for row in rows:
            target.append(row)
        for cell in target[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        target.freeze_panes = "A2"
        target.auto_filter.ref = target.dimensions
        for index, header in enumerate(headers, start=1):
            widths = [len(header) * 2]
            widths.extend(
                min(50, len(str(row[index - 1] or "")) + 2)
                for row in rows[:200]
            )
            target.column_dimensions[get_column_letter(index)].width = max(12, max(widths))
        for row in target.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    summary_rows = []
    generated_extra = []
    historical_extra = []
    account_differences = []
    for item in full.get("comparisons") or []:
        summary_rows.append([
            item["合同号_1"],
            item.get("导入文件", ""),
            "是" if item["会计字段一致"] else "否",
            "是" if item["忽略顺序及暂估进项科目后一致"] else "否",
            "；".join(item["剩余不一致原因"]),
            "是" if item["科目汇总一致"] else "否",
            "是" if item["行顺序一致"] else "否",
            "是" if item["包含摘要完全一致"] else "否",
            "是" if item["存在可匹配分录"] else "否",
            item["完全一致分录数"],
            item["同科目借方/贷方及辅助核算汇总匹配数"],
            item["同科目借贷总额匹配数"],
            item["生成行数"],
            item["历史行数"],
            "；".join(item["历史凭证ID"]),
        ])
        for line in item["生成多出行"]:
            generated_extra.append([
                item["合同号_1"], line["科目编码"], line["借方"], line["贷方"],
                line["客户编码"], line["供应商编码"], line["存货档案"],
                line["数量"], line["单价"],
            ])
        for line in item["历史多出行"]:
            historical_extra.append([
                item["合同号_1"], line["科目编码"], line["借方"], line["贷方"],
                line["客户编码"], line["供应商编码"], line["存货档案"],
                line["数量"], line["单价"],
            ])
        for difference in item["科目汇总差异"]:
            account_differences.append([
                item["合同号_1"], difference["科目编码"],
                difference["生成借方"], difference["历史借方"],
                difference["生成贷方"], difference["历史贷方"],
            ])

    add_sheet("差异汇总", [
        "合同号_1", "导入文件", "严格会计字段一致", "忽略顺序及暂估进项科目后一致",
        "剩余不一致原因", "科目汇总一致", "行顺序一致",
        "包含摘要完全一致", "存在可匹配分录", "完全一致分录数",
        "维度汇总匹配数", "科目汇总匹配数", "生成行数", "历史行数", "历史凭证ID",
    ], summary_rows)
    add_sheet("生成多出行", [
        "合同号_1", "科目编码", "借方", "贷方", "客户编码",
        "供应商编码", "存货档案", "数量", "单价",
    ], generated_extra)
    add_sheet("历史多出行", [
        "合同号_1", "科目编码", "借方", "贷方", "客户编码",
        "供应商编码", "存货档案", "数量", "单价",
    ], historical_extra)
    add_sheet("科目汇总差异", [
        "合同号_1", "科目编码", "生成借方", "历史借方", "生成贷方", "历史贷方",
    ], account_differences)

    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".part")
    workbook.save(temporary)
    temporary.replace(path)

